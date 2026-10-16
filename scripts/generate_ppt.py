import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Color Palette Definitions
GREEN_RGB = RGBColor(111, 172, 70)
YELLOW_RGB = RGBColor(255, 192, 0)
BLACK_RGB = RGBColor(50, 50, 50)  # Off-black for better typography readability
WHITE_RGB = RGBColor(255, 255, 255)
LIGHT_GREEN_RGB = RGBColor(235, 245, 235)
GRAY_RGB = RGBColor(100, 100, 100)  # Muted grey for sub-points

def apply_template(slide, prs):
    # 1. Top Green Bar
    top_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0), Inches(0), prs.slide_width, Inches(0.18)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = GREEN_RGB
    top_bar.line.fill.background()

    # 2. Bottom Yellow Bar
    bottom_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0), prs.slide_height - Inches(0.18), prs.slide_width, Inches(0.18)
    )
    bottom_bar.fill.solid()
    bottom_bar.fill.fore_color.rgb = YELLOW_RGB
    bottom_bar.line.fill.background()

    # 3. Logo (Top Right) - Perfect 1:1 Aspect Ratio (0.95" x 0.95")
    logo_path = r"C:\Users\anand\.gemini\antigravity\brain\75352bd4-1d28-4746-b2c3-bccbd8897f90\scratch\extracted_images\page_2_img_1.jpg"
    if os.path.exists(logo_path):
        slide.shapes.add_picture(
            logo_path, 
            prs.slide_width - Inches(1.35), Inches(0.3), 
            width=Inches(0.95), height=Inches(0.95)
        )

def add_full_image_slide(prs, image_path):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    if os.path.exists(image_path):
        slide.shapes.add_picture(
            image_path, 
            Inches(0), Inches(0), 
            width=prs.slide_width, height=prs.slide_height
        )
    return slide

def add_content_slide(prs, title_text, bullet_points):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_template(slide, prs)

    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.65), Inches(11.83), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = 'Arial'
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = GREEN_RGB

    # Content Box
    content_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.7), Inches(11.83), Inches(5.0))
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    tf_content.margin_left = tf_content.margin_right = tf_content.margin_top = tf_content.margin_bottom = 0

    first_p = True
    for pt in bullet_points:
        if isinstance(pt, tuple):
            main_pt, sub_pts = pt
            # Add main bullet (prepended with bullet dot)
            p = tf_content.paragraphs[0] if first_p else tf_content.add_paragraph()
            first_p = False
            p.text = f"•  {main_pt}"
            p.font.name = 'Arial'
            p.font.size = Pt(20)
            p.font.bold = True
            p.font.color.rgb = BLACK_RGB
            p.space_before = Pt(12)
            p.space_after = Pt(6)
            p.level = 0
            
            for sub_pt in sub_pts:
                sp = tf_content.add_paragraph()
                sp.text = f"    -  {sub_pt}"
                sp.font.name = 'Arial'
                sp.font.size = Pt(15)
                sp.font.color.rgb = GRAY_RGB
                sp.space_after = Pt(4)
                sp.level = 1
        else:
            p = tf_content.paragraphs[0] if first_p else tf_content.add_paragraph()
            first_p = False
            p.text = f"•  {pt}"
            p.font.name = 'Arial'
            p.font.size = Pt(18)
            p.font.color.rgb = BLACK_RGB
            p.space_before = Pt(8)
            p.space_after = Pt(8)
            p.level = 0

    return slide

def add_split_slide(prs, title_text, bullet_points, image_path):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_template(slide, prs)

    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.65), Inches(11.83), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = 'Arial'
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = GREEN_RGB

    # Left Column (Content)
    content_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.7), Inches(5.8), Inches(5.0))
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    tf_content.margin_left = tf_content.margin_right = tf_content.margin_top = tf_content.margin_bottom = 0

    first_p = True
    for pt in bullet_points:
        if isinstance(pt, tuple):
            main_pt, sub_pts = pt
            p = tf_content.paragraphs[0] if first_p else tf_content.add_paragraph()
            first_p = False
            p.text = f"•  {main_pt}"
            p.font.name = 'Arial'
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = BLACK_RGB
            p.space_before = Pt(10)
            p.space_after = Pt(6)
            p.level = 0
            
            for sub_pt in sub_pts:
                sp = tf_content.add_paragraph()
                sp.text = f"    -  {sub_pt}"
                sp.font.name = 'Arial'
                sp.font.size = Pt(14)
                sp.font.color.rgb = GRAY_RGB
                sp.space_after = Pt(4)
                sp.level = 1
        else:
            p = tf_content.paragraphs[0] if first_p else tf_content.add_paragraph()
            first_p = False
            p.text = f"•  {pt}"
            p.font.name = 'Arial'
            p.font.size = Pt(16)
            p.font.color.rgb = BLACK_RGB
            p.space_before = Pt(8)
            p.space_after = Pt(8)
            p.level = 0

    # Right Column (Image) - Perfect 16:9 Aspect Ratio (5.8" x 3.26")
    # Vertically centered in the remaining height (top=2.2")
    if os.path.exists(image_path):
        slide.shapes.add_picture(
            image_path, 
            Inches(6.8), Inches(2.2), 
            width=Inches(5.8), height=Inches(3.26)
        )

    return slide

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Slide 1: Cover Page (Collage)
    cover_img = r"C:\Users\anand\.gemini\antigravity\brain\75352bd4-1d28-4746-b2c3-bccbd8897f90\scratch\page_1_render.png"
    add_full_image_slide(prs, cover_img)

    # 2. Slide 2: Project Cover & Member Details
    blank_layout = prs.slide_layouts[6]
    title_slide = prs.slides.add_slide(blank_layout)
    apply_template(title_slide, prs)

    # "Industrial Training" Title
    it_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(0.7))
    tf_it = it_box.text_frame
    p_it = tf_it.paragraphs[0]
    p_it.text = "Industrial Training"
    p_it.font.name = 'Arial'
    p_it.font.size = Pt(32)
    p_it.font.color.rgb = BLACK_RGB
    p_it.alignment = PP_ALIGN.CENTER

    # Project Title
    pt_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.333), Inches(0.9))
    tf_pt = pt_box.text_frame
    p_pt = tf_pt.paragraphs[0]
    p_pt.text = "Visitor Management System"
    p_pt.font.name = 'Arial'
    p_pt.font.size = Pt(44)
    p_pt.font.bold = True
    p_pt.font.color.rgb = GREEN_RGB
    p_pt.alignment = PP_ALIGN.CENTER

    # Member Table
    table_shape = title_slide.shapes.add_table(2, 3, Inches(2.2), Inches(3.2), Inches(8.933), Inches(1.0))
    table = table_shape.table
    table.columns[0].width = Inches(2.8)
    table.columns[1].width = Inches(2.8)
    table.columns[2].width = Inches(3.333)

    # Headers
    headers = ["Member's Name", "Roll Number", "Branch"]
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = GREEN_RGB
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.name = 'Arial'
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE_RGB
        p.alignment = PP_ALIGN.CENTER

    # Data
    data = ["Amit Kumar", "1/24/SET/BCS/179", "Computer Science & Engineering"]
    for col_idx, text in enumerate(data):
        cell = table.cell(1, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = LIGHT_GREEN_RGB
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.name = 'Arial'
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = BLACK_RGB
        p.alignment = PP_ALIGN.CENTER

    # Guidance Details
    gd_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.333), Inches(1.2))
    tf_gd = gd_box.text_frame
    
    p_ug = tf_gd.paragraphs[0]
    p_ug.text = "Under the guidance of"
    p_ug.font.name = 'Arial'
    p_ug.font.size = Pt(18)
    p_ug.font.bold = True
    p_ug.font.underline = True
    p_ug.font.color.rgb = BLACK_RGB
    p_ug.alignment = PP_ALIGN.CENTER
    p_ug.space_after = Pt(4)

    p_mn = tf_gd.add_paragraph()
    p_mn.text = "Dr. Payal Thakur"
    p_mn.font.name = 'Arial'
    p_mn.font.size = Pt(22)
    p_mn.font.bold = True
    p_mn.font.color.rgb = GREEN_RGB
    p_mn.alignment = PP_ALIGN.CENTER

    p_title_dept = tf_gd.add_paragraph()
    p_title_dept.text = "Assistant Professor, Department of CSE"
    p_title_dept.font.name = 'Arial'
    p_title_dept.font.size = Pt(14)
    p_title_dept.font.color.rgb = BLACK_RGB
    p_title_dept.alignment = PP_ALIGN.CENTER

    # 3. Slide 3: Table of Contents (Exact Wording from PDF)
    contents_slide = prs.slides.add_slide(blank_layout)
    apply_template(contents_slide, prs)

    title_box = contents_slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.83), Inches(0.8))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = "CONTENTS"
    p_title.font.name = 'Arial'
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.underline = True
    p_title.font.color.rgb = BLACK_RGB
    p_title.alignment = PP_ALIGN.CENTER

    # Left Column Contents
    col1_box = contents_slide.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(5.2), Inches(5.0))
    tf_col1 = col1_box.text_frame
    col1_points = [
        "1. Introduction",
        "2. Company Details/ STIPHEND Details",
        "3. Objectives of Internship",
        "4. Project /Task Overview/ Problem Statement",
        "5. Technical Skills Acquired/ to be acquired",
        "6. Challenges Faced",
        "7. Achievements and Contributions"
    ]
    for i, pt in enumerate(col1_points):
        p = tf_col1.add_paragraph() if i > 0 else tf_col1.paragraphs[0]
        p.text = pt
        p.font.name = 'Arial'
        p.font.size = Pt(18)
        p.font.color.rgb = BLACK_RGB
        p.space_after = Pt(14)

    # Right Column Contents
    col2_box = contents_slide.shapes.add_textbox(Inches(7.2), Inches(1.6), Inches(5.2), Inches(5.0))
    tf_col2 = col2_box.text_frame
    col2_points = [
        "8. Learning and Insights",
        "9. Industry Trends and Observations",
        "10. Project Modules",
        "11. Methodology Used",
        "12. Result/Output (till now with screenshots)",
        "13. Conclusion and future scope",
        "14. References"
    ]
    for i, pt in enumerate(col2_points):
        p = tf_col2.add_paragraph() if i > 0 else tf_col2.paragraphs[0]
        p.text = pt
        p.font.name = 'Arial'
        p.font.size = Pt(18)
        p.font.color.rgb = BLACK_RGB
        p.space_after = Pt(14)

    # 4. Slide 4: Introduction
    intro_pts = [
        ("Digital Lobby Management", [
            "Developed a web-based, real-time Visitor Management System (VMS) to secure and streamline visitor check-in workflows.",
            "Automates the entire visitor lifecycle from self-registration to host approval, security verification, and check-out."
        ]),
        ("Shortcomings of Traditional Paper Logbooks", [
            "Lobby Congestion: Slow manual transcribing causes reception desk bottlenecks.",
            "No ID Verification: Security staff cannot verify uploaded identity documents for authenticity.",
            "Communication Delays: Guards must manually phone employees, slowing down host notifications.",
            "Data Privacy Violations: Leaving physical logbooks open on desks exposes visitor records to other guests."
        ])
    ]
    add_content_slide(prs, "1. Introduction", intro_pts)

    # 5. Slide 5: Company & Stipend Details
    company_pts = [
        ("Company Profile: Groz Tools Private Limited", [
            "Established in 1976 by Anil Bammi; started as a bicycle component manufacturer.",
            "Now a global leader in precision tooling, lubrication equipment, and hand hardware.",
            "Facilities span 500,000+ sq. ft. across Gurgaon, Noida, Chennai, with HQ in Gurgaon, Haryana.",
            "Employs 2,500+ professionals and exports 20,000+ products to 85+ countries."
        ]),
        ("Internship & STIPHEND Details", [
            "Duration: 4-week software engineering internship (June 25, 2025 – July 25, 2025).",
            "Role & Team: Software Engineering division working on internal enterprise Web Systems.",
            "Stipend Structure: Non-financial, educational internship focused on core skills development, enterprise architecture, and hands-on system building under professional industry mentors."
        ])
    ]
    add_content_slide(prs, "2. Company Details/ STIPHEND Details", company_pts)

    # 6. Slide 6: Objectives of Internship
    obj_pts = [
        ("Enterprise Software Development Lifecycle (SDLC)", [
            "Learn professional software development from schema modeling to production deployment."
        ]),
        ("Advanced Backend Architectures", [
            "Master event-driven, asynchronous Node.js & Express servers.",
            "Implement security practices (Helmet headers, rate limits, HTTPOnly JWT cookies)."
        ]),
        ("Real-Time Notification Systems", [
            "Establish full-duplex communication pipelines using Socket.io WebSockets for instant host alerts."
        ]),
        ("Artificial Intelligence & Computer Vision Integration", [
            "Integrate Google Gemini AI multimodal API to automate text extraction from identity card scans, reducing data entry errors."
        ])
    ]
    add_content_slide(prs, "3. Objectives of Internship", obj_pts)

    # 7. Slide 7: Project Overview & Problem Statement
    prob_pts = [
        ("Problem Statement", [
            "Traditional visitor tracking relies on manual paper logbooks, which are inefficient, slow, prone to data forgery, and pose privacy compliance risks under modern standards."
        ]),
        ("Project Overview", [
            "The Visitor Management System (VMS) is a centralized web platform digitizing physical access management.",
            "It integrates a Visitor Portal (registration), Host Employee Dashboard (approval), and a Security Guard Terminal (verification)."
        ]),
        ("Compliance & Scope", [
            "Includes automated email passes, QR code generation, real-time alerts, and auto-purge databases in compliance with data privacy policies."
        ])
    ]
    add_content_slide(prs, "4. Project /Task Overview/ Problem Statement", prob_pts)

    # 8. Slide 8: Technical Skills Acquired
    skills_pts = [
        ("Backend & Database Technologies", [
            "Node.js & Express: Asynchronous runtime and routing frameworks.",
            "MySQL: Relational database design, connection pooling, and self-healing schema migrations."
        ]),
        ("Real-Time & Communications", [
            "Socket.io: WebSockets for persistent client-server broadcast pipelines.",
            "Nodemailer / Brevo SMTP: Sending automated email alerts and PDF-like passes."
        ]),
        ("Cloud & AI Integration", [
            "Google Gemini AI API: Multimodal model for image OCR document analysis.",
            "Cloudinary CDN: Cloud storage and optimization for user-uploaded images and ID scans."
        ]),
        ("Security & Middleware", [
            "Helmet (secure headers), express-rate-limit (DDoS prevention), bcryptjs (password hashing), and JWT (access token management)."
        ])
    ]
    add_content_slide(prs, "5. Technical Skills Acquired/ to be acquired", skills_pts)

    # 9. Slide 9: Challenges Faced
    challenges_pts = [
        ("Database Pool Management", [
            "Challenge: Handling multiple concurrent queries caused database timeout and lock errors.",
            "Solution: Implemented MySQL Connection Pooling (limit of 10 connections) with automatic error fallbacks."
        ]),
        ("Cloud Upload Pipelines", [
            "Challenge: Large photo uploads blocked backend event loops and consumed disk space.",
            "Solution: Utilized Multer memory storage and Cloudinary SDK to stream uploads directly to the cloud CDN."
        ]),
        ("Real-Time Connection Stability", [
            "Challenge: Socket disconnections behind corporate firewalls.",
            "Solution: Configured Socket.io automatic reconnection and HTTP long-polling fallbacks."
        ]),
        ("Multimodal ID Verification", [
            "Challenge: Hand-held visitor ID scans were often low-resolution, tilted, or had poor lighting.",
            "Solution: Engineered detailed prompts with Gemini 2.5/1.5 Flash to extract text and return structured JSON."
        ])
    ]
    add_content_slide(prs, "6. Challenges Faced", challenges_pts)

    # 10. Slide 10: Achievements and Contributions
    ach_pts = [
        ("Full-Stack Implementation", [
            "Developed all three frontend dashboards and integrated them with a unified Express backend."
        ]),
        ("AI OCR Verification System", [
            "Designed and coded the `/ocr/upload` pipeline that automatically parses identity document images and compares details against visitor entries."
        ]),
        ("Automated Retention Purge Service", [
            "Built a scheduler running in the background to automatically delete visitor logs and their Cloudinary assets older than 30 days, ensuring privacy compliance."
        ]),
        ("Secure Cookie-Based JWT Auth", [
            "Configured HTTPOnly and Secure cookie rotation to manage guard and employee sessions safely."
        ])
    ]
    add_content_slide(prs, "7. Achievements and Contributions", ach_pts)

    # 11. Slide 11: Learning and Insights
    learning_pts = [
        ("Bridging Academic and Industry Practice", [
            "Classrooms teach theoretical databases and networking; the internship taught real-world limits, CORS restrictions, rate limit blockages, and asynchronous callback patterns."
        ]),
        ("Security-First Design Principles", [
            "Learned that data privacy is critical. Implementing XSS headers and encrypting session tokens is just as important as core functionality."
        ]),
        ("Modular Programming Advantages", [
            "Writing modular functions (separating routes, controllers, DB configs, and services) simplifies code maintenance and makes unit testing possible."
        ]),
        ("Version Control & Collaborative Development", [
            "Gained hands-on experience using Git to manage features, branches, and resolve conflicts."
        ])
    ]
    add_content_slide(prs, "8. Learning and Insights", learning_pts)

    # 12. Slide 12: Industry Trends & Observations
    trends_pts = [
        ("Touchless & Mobile Registration", [
            "Lobbies are moving away from centralized terminals. Modern systems let visitors register on their own mobile devices by scanning a simple lobby QR code."
        ]),
        ("Cloud-Native Infrastructure", [
            "Scaling image storage and asset pipelines dynamically using CDNs (Cloudinary, AWS S3) is standard practice in modern enterprise software."
        ]),
        ("AI Automation in Administration", [
            "Using AI (like Gemini) to read documents or match names is replacing manual data entry, improving security audits while reducing lobby wait times."
        ]),
        ("Global Data Privacy Compliance", [
            "Data privacy laws (GDPR, CCPA, local data bills) mandate strict data retention bounds, making automatic purging a necessary architectural requirement."
        ])
    ]
    add_content_slide(prs, "9. Industry Trends and Observations", trends_pts)

    # Images Paths for Split Layout Slides
    img_vreg = r"C:\Users\anand\.gemini\antigravity\brain\75352bd4-1d28-4746-b2c3-bccbd8897f90\visitor_registration_ui_1787112877828.jpg"
    img_host = r"C:\Users\anand\.gemini\antigravity\brain\75352bd4-1d28-4746-b2c3-bccbd8897f90\employee_dashboard_ui_1787112938401.jpg"
    img_sec = r"C:\Users\anand\.gemini\antigravity\brain\75352bd4-1d28-4746-b2c3-bccbd8897f90\security_guard_ui_1787113009239.jpg"
    img_arch = r"C:\Users\anand\.gemini\antigravity\brain\75352bd4-1d28-4746-b2c3-bccbd8897f90\system_architecture_diagram_1787113028020.jpg"

    # 13. Slide 13: Project Modules - Visitor Self-Registration
    mod1_pts = [
        ("Module 1: Visitor Portal Features", [
            "Self-onboarding: Visitors scan a lobby QR code and register on their personal mobile devices.",
            "Form inputs: Full Name, Phone Number, Company, Reason, Host Name, and Department.",
            "Webcam capture: Captures the visitor's photograph using native browser APIs.",
            "ID Upload: Multiparts document upload (Adhaar, license, etc.).",
            "Secure Pass: Backend uploads images to Cloudinary, creates a 'Pending' database log, generates a unique QR code pass, and emails it to the visitor."
        ])
    ]
    add_split_slide(prs, "10. Project Modules: Visitor Registration", mod1_pts, img_vreg)

    # 14. Slide 14: Project Modules - Host Portal & Real-Time Approval
    mod2_pts = [
        ("Module 2: Host Portal Features", [
            "Security-locked dashboard: Host employees authenticate via JWT cookies.",
            "Real-time notifications: Established WebSockets (Socket.io) to push check-in alerts instantly.",
            "Visual verification: Employee sees visitor photograph, name, and purpose.",
            "One-click actions: Employees click 'Approve' (which triggers automated email passes) or 'Reject' (with a reason, notifying the visitor)."
        ])
    ]
    add_split_slide(prs, "10. Project Modules: Host Approval Portal", mod2_pts, img_host)

    # 15. Slide 15: Project Modules - Security Guard Console & AI OCR
    mod3_pts = [
        ("Module 3: Security Console Features", [
            "QR Pass Scanner: Guards scan QR passes using webcam readers to query check-in status.",
            "System Stats: Displays total pending, checked-in, and checked-out visitors.",
            "AI OCR Audits: Uses Google Gemini multimodal API to read text from uploaded ID cards.",
            "Discrepancy Checks: Compares OCR name/number with database records, showing clear alerts if they mismatch."
        ])
    ]
    add_split_slide(prs, "10. Project Modules: Security & OCR Verification", mod3_pts, img_sec)

    # 16. Slide 16: Methodology & System Architecture
    meth_pts = [
        ("Agile Methodology", [
            "Development split into 2-week sprints with ongoing feedback loops and rapid prototyping."
        ]),
        ("System Architecture Design", [
            "Model-View-Controller backend separation.",
            "Client: Web Browser (HTML, CSS, JS, webcam, QR).",
            "Server: Node.js with Express framework.",
            "Database: MySQL Server with Connection Pooling.",
            "CDN Services: Cloudinary for fast image asset delivery.",
            "AI Layer: Google Gemini SDK for OCR document checks."
        ])
    ]
    add_split_slide(prs, "11. Methodology Used: Architecture", meth_pts, img_arch)

    # 17. Slide 17: Result/Output
    res_pts = [
        ("Self-Healing Schema Migrations", [
            "Backend automatically verifies and updates database columns during initialization, removing manual intervention."
        ]),
        ("Real-Time Alerts", [
            "Socket.io triggers immediate audio-visual alerts on host dashboards with zero lag."
        ]),
        ("AI-Powered Verification Panel", [
            "Gemini OCR correctly extracts text from ID uploads and compares it to user inputs, flagging mismatches."
        ]),
        ("Automatic Data Purging", [
            "Logs older than 30 days are automatically deleted from the MySQL database, and images are cleared from Cloudinary."
        ])
    ]
    add_content_slide(prs, "12. Result/Output (till now with screenshots)", res_pts)

    # 18. Slide 18: Conclusion and Future Scope
    conclusion_pts = [
        ("Conclusion", [
            "The VMS project successfully modernizes reception security, increases verification speed, reduces data entry errors, and satisfies corporate compliance policies at Groz Tools."
        ]),
        ("Future Scope", [
            "Biometric & Face Matching: Automatically compare webcam photos with identity document pictures.",
            "Hardware Integration: Connect the database to turnstile gates to allow automatic entry on valid QR scans.",
            "Active Directory Integration: Sync employee lists automatically with corporate LDAP/Azure AD."
        ])
    ]
    add_content_slide(prs, "13. Conclusion and future scope", conclusion_pts)

    # 19. Slide 19: References
    ref_pts = [
        "1. Node.js Documentation and API Reference - https://nodejs.org/api/",
        "2. ExpressJS Routing and Web Development Guide - https://expressjs.com/",
        "3. Socket.io WebSocket Communication Guide - https://socket.io/docs/v4/",
        "4. MySQL Connection Pooling and mysql2 Reference - https://sidorares.github.io/node-mysql2/",
        "5. Cloudinary Node.js Image Management SDK - https://cloudinary.com/documentation/",
        "6. Google Gemini AI Developer Documentation - https://ai.google.dev/",
        "7. Groz Tools Private Limited Corporate Website - https://www.groz-tools.com/"
    ]
    add_content_slide(prs, "14. References", ref_pts)

    # 20. Slide 20: Thank You Slide
    thank_img = r"C:\Users\anand\.gemini\antigravity\brain\75352bd4-1d28-4746-b2c3-bccbd8897f90\scratch\page_6_render.png"
    add_full_image_slide(prs, thank_img)

    # Save Presentation
    output_path = r"c:\Users\anand\OneDrive\Desktop\visitor-management\Visitor_Management_System_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    main()
